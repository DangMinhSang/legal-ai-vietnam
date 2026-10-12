"""Comprehensive document parser implementation for Legal AI Vietnam.

Supports PDF (text extraction with OCR fallback), DOCX, and TXT files.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from loguru import logger
from src.document_parser.base import BaseParser, ParsedDocument


class DocumentParser(BaseParser):
    """Universal parser for PDF, DOCX, and plain text files."""

    SUPPORTED_EXTENSIONS = frozenset({".txt", ".pdf", ".docx"})

    def __init__(self, enable_ocr: bool = True) -> None:
        self.enable_ocr = enable_ocr

    def supports(self, source: Path | str) -> bool:
        """Check if file extension is supported."""
        suffix = Path(source).suffix.lower()
        return suffix in self.SUPPORTED_EXTENSIONS

    def parse(self, source: Path | bytes | str) -> ParsedDocument:
        """Parse document and return normalised ParsedDocument."""
        if isinstance(source, bytes):
            text = source.decode("utf-8", errors="replace")
            cleaned = self._clean_text(text)
            if not cleaned:
                raise ValueError("File content is empty")
            return ParsedDocument(
                title=None,
                content=cleaned,
                page_count=1,
                parse_method="text_extraction",
            )

        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        suffix = path.suffix.lower()
        if suffix not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file format: '{suffix}'. Supported: {sorted(self.SUPPORTED_EXTENSIONS)}")

        # Check if file is completely empty (0 bytes)
        if path.stat().st_size == 0:
            raise ValueError(f"File is empty: {path}")

        if suffix == ".txt":
            return self._parse_txt(path)
        elif suffix == ".docx":
            return self._parse_docx(path)
        elif suffix == ".pdf":
            return self._parse_pdf(path)

        raise ValueError(f"Unsupported file format: {suffix}")

    def _parse_txt(self, path: Path) -> ParsedDocument:
        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            content = path.read_text(encoding="latin-1", errors="replace")

        cleaned = self._clean_text(content)
        if not cleaned:
            raise ValueError(f"Text file is empty or contains only whitespace: {path}")

        return ParsedDocument(
            title=path.stem,
            content=cleaned,
            source_path=str(path),
            page_count=1,
            parse_method="text_extraction",
        )

    def _parse_docx(self, path: Path) -> ParsedDocument:
        try:
            import docx
            doc = docx.Document(path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n\n".join(paragraphs)
            cleaned = self._clean_text(content)
            if not cleaned:
                raise ValueError(f"DOCX file contains no readable text: {path}")

            title = paragraphs[0] if paragraphs else path.stem
            return ParsedDocument(
                title=title[:200],
                content=cleaned,
                source_path=str(path),
                page_count=1,
                parse_method="text_extraction",
            )
        except ImportError:
            raise RuntimeError("python-docx is required to parse .docx files")

    def _parse_pdf(self, path: Path) -> ParsedDocument:
        """Parse PDF: extract text layer first; if empty and OCR is enabled, fallback to OCR."""
        text_pages: list[str] = []
        page_count = 0
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            page_count = len(reader.pages)
            for page in reader.pages:
                extracted = page.extract_text() or ""
                if extracted.strip():
                    text_pages.append(extracted)
        except Exception as exc:
            logger.warning("Error reading PDF text layer: {}", exc)

        full_text = "\n\n".join(text_pages).strip()
        cleaned = self._clean_text(full_text)

        # If text layer found, PREFER text layer and DO NOT run OCR
        if cleaned:
            return ParsedDocument(
                title=path.stem,
                content=cleaned,
                source_path=str(path),
                page_count=page_count,
                parse_method="text_extraction",
            )

        # Fallback to OCR if no text layer and OCR enabled
        if self.enable_ocr:
            ocr_text = self._run_ocr(path)
            cleaned_ocr = self._clean_text(ocr_text)
            if cleaned_ocr:
                return ParsedDocument(
                    title=path.stem,
                    content=cleaned_ocr,
                    source_path=str(path),
                    page_count=page_count,
                    parse_method="ocr",
                    warnings=["Text extracted via OCR fallback"],
                )

        raise ValueError(f"PDF file has no readable text and OCR produced no output: {path}")

    def _run_ocr(self, path: Path) -> str:
        """Run OCR using rapidocr if available."""
        try:
            from rapidocr_onnxruntime import RapidOCR
            engine = RapidOCR()
            result, _ = engine(str(path))
            if result:
                return "\n".join([line[1] for line in result])
        except Exception as exc:
            logger.warning("RapidOCR failed: {}", exc)
        return ""

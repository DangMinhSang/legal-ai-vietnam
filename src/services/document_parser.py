"""Document parser service supporting PDF, DOCX, and TXT formats with OCR fallback.

Provides:
- ParsedDocument: Dataclass containing parsed text, per-page breakdown, metadata, and error details.
- DocumentParser: Unified document parser with automatic OCR fallback for scanned PDFs.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from loguru import logger
from pypdf import PdfReader
from docx import Document as DocxDocument
try:
    from rapidocr_onnxruntime import RapidOCR
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False


@dataclass
class ParsedDocument:
    """Standardized representation of a parsed document.

    Attributes:
        text: Full extracted or OCR-recognized text.
        pages: List of per-page dictionaries: {"page_num": int, "text": str, "method": str}.
        metadata: Document metadata (file name, size, page count, format details).
        parse_method: Method used: "text_extraction" | "ocr" | "mixed".
        confidence: Overall parsing confidence score between 0.0 and 1.0.
        errors: List of warnings, non-fatal errors, or failure descriptions.
    """

    text: str = ""
    pages: list[dict[str, Any]] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)
    parse_method: str = "text_extraction"
    confidence: float = 1.0
    errors: list[str] = field(default_factory=list)


class DocumentParser:
    """Production-ready document parser for legal texts in PDF, DOCX, and TXT formats.

    Features:
    - Text extraction first for PDF documents via pypdf.
    - Automatic fallback to OCR (rapidocr-onnxruntime) when text density < 100 chars/page.
    - Word document (.docx) parsing via python-docx.
    - Plain text (.txt) parsing with UTF-8 encoding detection.
    - Never raises unhandled exceptions; returns partial results with error diagnostics.
    """

    def __init__(self, ocr_engine: Any = None) -> None:
        """Initialize DocumentParser.

        Args:
            ocr_engine: Optional pre-configured RapidOCR instance.
        """
        self._ocr = ocr_engine

    def _get_ocr(self) -> Any:
        """Lazy-initialize and return the RapidOCR engine instance."""
        if self._ocr is None and OCR_AVAILABLE:
            try:
                self._ocr = RapidOCR()
            except Exception as e:
                logger.error(f"Failed to initialize RapidOCR engine: {e}")
                self._ocr = None
        return self._ocr

    def parse(self, file_path: Path | str) -> ParsedDocument:
        """Parse a document file and return standardized ParsedDocument.

        Args:
            file_path: Path to the target document file.

        Returns:
            ParsedDocument containing extracted text, pages, metadata, and status.
            Never raises exceptions.
        """
        try:
            path = Path(file_path) if not isinstance(file_path, Path) else file_path
        except Exception as e:
            logger.error(f"Invalid file path provided: {file_path!r} ({e})")
            return ParsedDocument(
                text="",
                pages=[],
                metadata={"file_path": str(file_path)},
                parse_method="text_extraction",
                confidence=0.0,
                errors=[f"Invalid file path: {str(e)}"],
            )

        if not path.exists():
            msg = f"File not found: {path}"
            logger.warning(msg)
            return ParsedDocument(
                text="",
                pages=[],
                metadata={"file_path": str(path)},
                parse_method="text_extraction",
                confidence=0.0,
                errors=[msg],
            )

        if path.is_dir():
            msg = f"Path is a directory, not a file: {path}"
            logger.warning(msg)
            return ParsedDocument(
                text="",
                pages=[],
                metadata={"file_path": str(path)},
                parse_method="text_extraction",
                confidence=0.0,
                errors=[msg],
            )

        suffix = path.suffix.lower()

        try:
            if suffix == ".pdf":
                return self._parse_pdf(path)
            elif suffix in (".docx", ".doc"):
                return self._parse_docx(path)
            elif suffix in (".txt", ".text", ".log", ".md"):
                return self._parse_txt(path)
            else:
                msg = f"Unsupported file format: {suffix}"
                logger.warning(f"{msg} for file {path.name}")
                return ParsedDocument(
                    text="",
                    pages=[],
                    metadata={"file_path": str(path), "extension": suffix},
                    parse_method="text_extraction",
                    confidence=0.0,
                    errors=[msg],
                )
        except Exception as e:
            logger.error(f"Unexpected error while parsing {path}: {e}")
            return ParsedDocument(
                text="",
                pages=[],
                metadata={"file_path": str(path), "extension": suffix},
                parse_method="text_extraction",
                confidence=0.0,
                errors=[f"Parsing failed: {str(e)}"],
            )

    def _parse_txt(self, path: Path) -> ParsedDocument:
        """Parse plain text file with UTF-8 encoding support."""
        errors: list[str] = []
        content = ""
        encoding_used = "utf-8"

        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            try:
                content = path.read_text(encoding="utf-8-sig")
                encoding_used = "utf-8-sig"
            except UnicodeDecodeError:
                content = path.read_text(encoding="utf-8", errors="replace")
                encoding_used = "utf-8 (replace)"
                errors.append("UTF-8 decoding errors encountered; invalid characters replaced.")
        except Exception as e:
            errors.append(f"Failed to read text file: {str(e)}")
            return ParsedDocument(
                text="",
                pages=[],
                metadata={"file_name": path.name, "file_path": str(path)},
                parse_method="text_extraction",
                confidence=0.0,
                errors=errors,
            )

        pages = [{"page_num": 1, "text": content, "method": "text_extraction"}]
        metadata = {
            "file_name": path.name,
            "file_path": str(path),
            "file_size": path.stat().st_size if path.exists() else 0,
            "total_pages": 1,
            "encoding": encoding_used,
            "extension": path.suffix.lower(),
        }

        return ParsedDocument(
            text=content,
            pages=pages,
            metadata=metadata,
            parse_method="text_extraction",
            confidence=1.0,
            errors=errors,
        )

    def _parse_docx(self, path: Path) -> ParsedDocument:
        """Parse Microsoft Word DOCX document."""
        errors: list[str] = []
        try:
            doc = DocxDocument(str(path))
            paragraph_lines: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]

            table_lines: list[str] = []
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        table_lines.append(" | ".join(row_cells))

            full_text = "\n".join(paragraph_lines + table_lines)
            pages = [{"page_num": 1, "text": full_text, "method": "text_extraction"}]
            metadata = {
                "file_name": path.name,
                "file_path": str(path),
                "file_size": path.stat().st_size if path.exists() else 0,
                "total_pages": 1,
                "paragraph_count": len(paragraph_lines),
                "table_count": len(doc.tables),
                "extension": path.suffix.lower(),
            }

            return ParsedDocument(
                text=full_text,
                pages=pages,
                metadata=metadata,
                parse_method="text_extraction",
                confidence=1.0,
                errors=errors,
            )
        except Exception as e:
            logger.error(f"DOCX extraction failed for {path}: {e}")
            errors.append(f"DOCX extraction failed: {str(e)}")
            return ParsedDocument(
                text="",
                pages=[],
                metadata={"file_name": path.name, "file_path": str(path)},
                parse_method="text_extraction",
                confidence=0.0,
                errors=errors,
            )

    def _parse_pdf(self, path: Path) -> ParsedDocument:
        """Parse PDF document with text extraction and automatic OCR fallback."""
        import io
        errors: list[str] = []
        try:
            pdf_bytes = path.read_bytes()
            reader = PdfReader(io.BytesIO(pdf_bytes))
        except Exception as e:
            logger.error(f"Failed to open PDF {path}: {e}")
            errors.append(f"Failed to read PDF file: {str(e)}")
            return ParsedDocument(
                text="",
                pages=[],
                metadata={"file_name": path.name, "file_path": str(path)},
                parse_method="text_extraction",
                confidence=0.0,
                errors=errors,
            )

        num_pages = len(reader.pages)
        if num_pages == 0:
            return ParsedDocument(
                text="",
                pages=[],
                metadata={"file_name": path.name, "file_path": str(path), "total_pages": 0},
                parse_method="text_extraction",
                confidence=0.0,
                errors=["PDF contains 0 pages."],
            )

        metadata = {
            "file_name": path.name,
            "file_path": str(path),
            "file_size": path.stat().st_size if path.exists() else 0,
            "total_pages": num_pages,
            "extension": ".pdf",
        }

        # Step 1: Try text extraction first
        extracted_pages: list[dict[str, Any]] = []
        for idx, page in enumerate(reader.pages, start=1):
            page_text = ""
            try:
                page_text = page.extract_text() or ""
            except Exception as e:
                logger.warning(f"Error extracting text from page {idx} in {path.name}: {e}")
                errors.append(f"Page {idx} text extraction error: {str(e)}")

            extracted_pages.append({
                "page_num": idx,
                "text": page_text,
                "method": "text_extraction",
            })

        total_extracted_chars = sum(len(p["text"].strip()) for p in extracted_pages)
        avg_chars_per_page = total_extracted_chars / num_pages if num_pages > 0 else 0
        metadata["avg_chars_per_page"] = round(avg_chars_per_page, 2)

        # Step 2: If text density is sufficient (>= 100 chars/page average), return text_extraction
        if avg_chars_per_page >= 100:
            full_text = "\n\n".join(p["text"].strip() for p in extracted_pages if p["text"].strip())
            return ParsedDocument(
                text=full_text,
                pages=extracted_pages,
                metadata=metadata,
                parse_method="text_extraction",
                confidence=1.0,
                errors=errors,
            )

        # Step 3: Text is too short (< 100 chars/page average) -> fallback to OCR
        logger.info(
            f"PDF {path.name} average characters ({avg_chars_per_page:.1f}/page) < 100. "
            f"Falling back to OCR..."
        )

        if not OCR_AVAILABLE:
            msg = (
                f"PDF text density ({avg_chars_per_page:.1f} chars/page) < 100, "
                "but rapidocr-onnxruntime is not installed. Keeping extracted text."
            )
            logger.warning(msg)
            errors.append(msg)
            full_text = "\n\n".join(p["text"].strip() for p in extracted_pages if p["text"].strip())
            return ParsedDocument(
                text=full_text,
                pages=extracted_pages,
                metadata=metadata,
                parse_method="text_extraction",
                confidence=1.0,
                errors=errors,
            )

        ocr_engine = self._get_ocr()
        if ocr_engine is None:
            msg = "OCR fallback requested, but RapidOCR engine could not be initialized."
            logger.warning(msg)
            errors.append(msg)
            full_text = "\n\n".join(p["text"].strip() for p in extracted_pages if p["text"].strip())
            return ParsedDocument(
                text=full_text,
                pages=extracted_pages,
                metadata=metadata,
                parse_method="text_extraction",
                confidence=1.0,
                errors=errors,
            )

        final_pages: list[dict[str, Any]] = []
        page_confidences: list[float] = []

        for idx, page_info in enumerate(extracted_pages, start=1):
            curr_text = page_info["text"].strip()
            # If the page already has sufficient text (>= 100 chars), preserve text extraction
            if len(curr_text) >= 100:
                final_pages.append(page_info)
                page_confidences.append(1.0)
                continue

            # Fall back to OCR for pages with insufficient text
            pypdf_page = reader.pages[idx - 1]
            ocr_text, ocr_scores, ocr_err = self._ocr_pdf_page(path, idx - 1, pypdf_page, ocr_engine)
            if ocr_err:
                errors.append(ocr_err)

            if ocr_text.strip():
                final_pages.append({
                    "page_num": idx,
                    "text": ocr_text,
                    "method": "ocr",
                })
                avg_score = sum(ocr_scores) / len(ocr_scores) if ocr_scores else 0.5
                page_confidences.append(avg_score)
            else:
                # OCR yielded nothing; preserve whatever extracted text was present
                final_pages.append({
                    "page_num": idx,
                    "text": curr_text,
                    "method": "text_extraction",
                })
                page_confidences.append(1.0 if curr_text else 0.0)

        # Determine overall parse_method: "text_extraction" | "ocr" | "mixed"
        methods = {p["method"] for p in final_pages}
        if len(methods) > 1:
            parse_method = "mixed"
        elif "ocr" in methods:
            parse_method = "ocr"
        else:
            parse_method = "text_extraction"

        overall_confidence = (
            sum(page_confidences) / len(page_confidences) if page_confidences else 1.0
        )
        if overall_confidence < 0.7:
            logger.warning(
                f"Low overall OCR confidence for {path.name}: {overall_confidence:.2f} < 0.7"
            )

        full_text = "\n\n".join(p["text"].strip() for p in final_pages if p["text"].strip())
        return ParsedDocument(
            text=full_text,
            pages=final_pages,
            metadata=metadata,
            parse_method=parse_method,
            confidence=round(overall_confidence, 4),
            errors=errors,
        )

    def _ocr_pdf_page(
        self,
        file_path: Path,
        page_idx: int,
        pypdf_page: Any,
        ocr_engine: Any,
    ) -> tuple[str, list[float], str | None]:
        """Perform OCR on a single PDF page image."""
        page_num = page_idx + 1
        img = self._get_page_image(file_path, page_idx, pypdf_page)
        if img is None:
            return "", [], f"Page {page_num}: Could not render or extract page image for OCR."

        try:
            raw_res, _ = ocr_engine(img)
        except Exception as e:
            logger.warning(f"RapidOCR call failed on page {page_num}: {e}")
            return "", [], f"Page {page_num}: OCR processing error: {str(e)}"

        texts, scores = self._process_ocr_output(raw_res, page_num)
        return "\n".join(texts), scores, None

    def _process_ocr_output(
        self,
        raw_res: Any,
        page_num: int,
    ) -> tuple[list[str], list[float]]:
        """Process OCR output items, handling list of (text, score) tuples or RapidOCR results.

        Logs a warning whenever OCR confidence < 0.7.
        """
        if not raw_res:
            return [], []

        texts: list[str] = []
        scores: list[float] = []

        for item in raw_res:
            txt = ""
            score = 1.0
            try:
                if isinstance(item, (list, tuple)):
                    if len(item) == 2:
                        # Direct (text, score) tuple
                        txt = str(item[0])
                        score = float(item[1])
                    elif len(item) >= 3:
                        # RapidOCR format: [box, text, score]
                        txt = str(item[1])
                        score = float(item[2])
                    elif len(item) == 1:
                        txt = str(item[0])
                        score = 1.0
                elif isinstance(item, str):
                    txt = item
                    score = 1.0
            except Exception as e:
                logger.warning(f"Error extracting text and score from OCR item: {e}")
                continue

            txt = txt.strip()
            if txt:
                texts.append(txt)
                scores.append(score)
                if score < 0.7:
                    logger.warning(
                        f"Low OCR confidence ({score:.2f} < 0.7) on page {page_num} "
                        f"for text: {txt[:60]!r}"
                    )

        return texts, scores

    def _get_page_image(self, file_path: Path, page_idx: int, pypdf_page: Any) -> Any:
        """Render or extract an image for the given PDF page."""
        # Method 1: pypdfium2 (renders vector and raster elements at 2x scale)
        try:
            import pypdfium2 as pdfium
            pdf_bytes = file_path.read_bytes()
            pdf_doc = pdfium.PdfDocument(pdf_bytes)
            try:
                pdf_page = pdf_doc.get_page(page_idx)
                try:
                    bitmap = pdf_page.render(scale=2.0)
                    try:
                        arr = bitmap.to_numpy().copy()
                        return arr
                    finally:
                        bitmap.close()
                finally:
                    pdf_page.close()
            finally:
                pdf_doc.close()
        except Exception as e:
            logger.debug(f"pypdfium2 rendering failed for page {page_idx + 1}: {e}")

        # Method 2: Extract embedded raster images via pypdf
        try:
            if hasattr(pypdf_page, "images") and len(pypdf_page.images) > 0:
                import io
                from PIL import Image
                import numpy as np
                first_img = pypdf_page.images[0]
                pil_img = Image.open(io.BytesIO(first_img.data)).convert("RGB")
                return np.array(pil_img)
        except Exception as e:
            logger.debug(f"pypdf image extraction failed for page {page_idx + 1}: {e}")

        # Method 3: PyMuPDF / fitz fallback if available
        try:
            import fitz
            import io
            from PIL import Image
            import numpy as np
            doc = fitz.open(str(file_path))
            page = doc.load_page(page_idx)
            pix = page.get_pixmap(dpi=150)
            pil_img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
            return np.array(pil_img)
        except Exception as e:
            logger.debug(f"fitz rendering failed for page {page_idx + 1}: {e}")

        return None
